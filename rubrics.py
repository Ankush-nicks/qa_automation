"""Reads the 'Updated Rubrics and Scoring' workbook and returns the parts of it
that matter for one instructor's weak categories: the sub-rubrics, their
definitions, and the score-band descriptions for the current and next band."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

# Which sheets to use for each instructor category and session type (first match wins,
# later sheets fill in any blank definitions).
LECTURE_SHEETS = {"TECH": ["Tech Sep26", "Tech"], "ENGLISH": ["English"], "MATH": ["Math"], "APTITUDE": ["Aptitude"]}
PRACTICE_SHEETS = {"ENGLISH": ["English Mentor Category", "Mentor Category"]}
DEFAULT_PRACTICE = ["Mentor Category"]
INVIGILATION_SHEETS = ["Invigilation Category"]
OTHER_SHEETS = ["Other Category"]

VALID_BANDS = {(0.0, 1.0), (1.5, 2.5), (3.0, 3.5), (4.0, 4.5), (5.0, 5.5), (6.0, 6.5), (7.0, 7.5), (8.0, 8.5),
               (9.0, 10.0), (9.0, 9.5)}
BAND_RE = re.compile(r"(\d{1,2}\.\d)\s*[–—-]\s*(\d{1,2}\.\d)")


def norm(s: str) -> str:
    s = str(s or "").lower()
    s = re.sub(r"\(.*?\)", " ", s)
    s = s.replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def clean(s: str, limit: int) -> str:
    s = re.sub(r"\*\*|\|", " ", str(s or ""))
    s = re.sub(r"\s+", " ", s).strip(" -:")
    return s if len(s) <= limit else s[:limit].rsplit(" ", 1)[0] + "…"


def split_bands(text: str) -> list[dict]:
    """Splits a 'Score Guidelines' cell into bands like {'lo':5.0,'hi':5.5,'text':...}."""
    text = str(text or "")
    hits = []
    for m in BAND_RE.finditer(text):
        lo, hi = float(m[1]), float(m[2])
        if (lo, hi) in VALID_BANDS:
            hits.append((m.start(), m.end(), lo, hi))
    bands, seen = [], set()
    for i, (s, e, lo, hi) in enumerate(hits):
        end = hits[i + 1][0] if i + 1 < len(hits) else len(text)
        body = clean(text[e:end], 450)
        if (lo, hi) in seen or len(body) < 15:
            continue
        seen.add((lo, hi))
        bands.append({"lo": lo, "hi": hi, "text": body})
    bands.sort(key=lambda b: b["lo"])
    return bands


class RubricBook:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.sheets: dict[str, dict[str, dict]] = {}
        wb = load_workbook(self.path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            hdr_i = next((i for i, r in enumerate(rows[:10])
                          if r and any(str(c or "").strip().lower() == "category name" for c in r)), None)
            if hdr_i is None:
                continue
            hdr = [str(c or "").strip().lower() for c in rows[hdr_i]]

            def col(*keys):
                for k in keys:
                    for i, h in enumerate(hdr):
                        if k in h:
                            return i
                return None

            ci, ri, di, gi, wi = (col("category name"), col("rubric name"), col("definition"),
                                  col("guideline", "scoring"), col("weig"))
            if ci is None or ri is None:
                continue
            cats: dict[str, dict] = {}
            current = None
            for r in rows[hdr_i + 1:]:
                if not r:
                    continue
                get = lambda i: (r[i] if i is not None and i < len(r) else None)  # noqa: E731
                c = get(ci)
                if c and str(c).strip():
                    current = str(c).strip()
                name = get(ri)
                if not current or not name or not str(name).strip():
                    continue
                entry = cats.setdefault(norm(current), {"name": current, "items": []})
                w = get(wi)
                entry["items"].append({"name": str(name).strip(),
                                       "definition": str(get(di) or "").strip(),
                                       "guidelines": str(get(gi) or "").strip(),
                                       "weight": w if isinstance(w, (int, float)) else None})
            self.sheets[ws.title] = cats
        wb.close()

    def sheet_list(self, inst_category: str, session_type: str) -> list[str]:
        c = (inst_category or "").upper()
        if session_type == "Lecture":
            return LECTURE_SHEETS.get(c, ["Tech"])
        if session_type == "Practice":
            return PRACTICE_SHEETS.get(c, DEFAULT_PRACTICE)
        if session_type == "Invigilation":
            return INVIGILATION_SHEETS
        return OTHER_SHEETS

    def context(self, inst_category: str, session_type: str, category_key: str, score: float,
                target: float = 6.0) -> dict | None:
        """Sub-rubrics for one category: the band the score sits in, the next band up and the target band."""
        key = norm(category_key)
        merged: dict[str, dict] = {}
        cat_name = None
        for sheet in self.sheet_list(inst_category, session_type):
            entry = (self.sheets.get(sheet) or {}).get(key)
            if not entry:
                continue
            cat_name = cat_name or entry["name"]
            for it in entry["items"]:
                m = merged.setdefault(norm(it["name"]), {"name": it["name"], "definition": "", "guidelines": "",
                                                         "weight": None})
                for f in ("definition", "guidelines"):
                    if not m[f] and it[f]:
                        m[f] = it[f]
                if m["weight"] is None:
                    m["weight"] = it["weight"]
        if not merged:
            return None
        subs = []
        for m in merged.values():
            bands = split_bands(m["guidelines"])
            cur = nxt = tgt = None
            for b in bands:
                if b["lo"] <= score + 1e-9:
                    cur = b
                if b["lo"] <= target + 1e-9:
                    tgt = b
            if bands:
                higher = [b for b in bands if cur is None or b["lo"] > cur["lo"]]
                nxt = higher[0] if higher else None
            subs.append({
                "name": m["name"],
                "weight_within_category": m["weight"],
                "definition": clean(m["definition"], 600),
                "band_now": f"{cur['lo']}–{cur['hi']}: {cur['text']}" if cur else None,
                "next_band": f"{nxt['lo']}–{nxt['hi']}: {nxt['text']}" if nxt else None,
                "target_band": f"{tgt['lo']}–{tgt['hi']}: {tgt['text']}" if (tgt and tgt is not nxt and tgt is not cur) else None,
                "guidelines_excerpt": None if bands else (clean(m["guidelines"], 600) or None),
            })
        return {"category": cat_name, "sub_rubrics": subs}
