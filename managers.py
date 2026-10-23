"""Reads the capability manager list (name -> email) and finds each instructor's
capability manager among the names in the export's instructor_manager_name column
(column F), so their email can be CC'd."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

NAME_COL = "capability manager name"
EMAIL_COL = "capability manager email"


def norm_name(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def load_managers(path: str | Path) -> dict[str, dict]:
    """Returns {normalised name: {"name", "email"}} from the first sheet that has
    'Capability Manager Name' and 'Capability Manager Email' columns."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for i, r in enumerate(rows[:10]):
                hdr = [norm_name(c) for c in (r or [])]
                if NAME_COL in hdr and EMAIL_COL in hdr:
                    ni, ei = hdr.index(NAME_COL), hdr.index(EMAIL_COL)
                    out = {}
                    for row in rows[i + 1:]:
                        name = row[ni] if row and ni < len(row) else None
                        email = row[ei] if row and ei < len(row) else None
                        if norm_name(name) and str(email or "").strip():
                            out[norm_name(name)] = {"name": str(name).strip(), "email": str(email).strip().lower()}
                    return out
    finally:
        wb.close()
    return {}


def match_managers(manager_names: str, managers: dict[str, dict]) -> list[dict]:
    """The capability managers among a comma-separated list of names, in the order listed."""
    found, seen = [], set()
    for n in str(manager_names or "").split(","):
        m = managers.get(norm_name(n))
        if m and m["email"] not in seen:
            seen.add(m["email"])
            found.append(m)
    return found
